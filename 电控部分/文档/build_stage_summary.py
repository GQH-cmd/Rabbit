from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"E:\养兔子\自动化幼兔喂奶系统电控阶段性工作总结.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "666666"
BODY = "1F1F1F"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths_dxa, table_width=9360, indent=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(table_width))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_east_asia_font(run, east_asia="Microsoft YaHei", ascii_font="Calibri"):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def style_run(run, size=11, bold=False, color=BODY, italic=False):
    set_east_asia_font(run)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    style_run(run, size=9, color=MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Pt(22)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        style_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        style_run(r2)
    else:
        r = p.add_run(text)
        style_run(r)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    style_run(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    style_run(r)
    return p


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Pt(10)
    p.paragraph_format.right_indent = Pt(10)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_BLUE)
    p_pr.append(shd)
    r1 = p.add_run(label + "：")
    style_run(r1, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    style_run(r2)
    return p


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def build():
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header_p = section.header.paragraphs[0]
    header_p.text = "自动化幼兔喂奶与清洗控制系统｜电控阶段性总结"
    header_p.paragraph_format.space_after = Pt(0)
    for run in header_p.runs:
        style_run(run, size=9, color=MUTED)

    footer_p = section.footer.paragraphs[0]
    add_page_number(footer_p)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("阶段性工作总结")
    style_run(r, size=24, bold=True, color="000000")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("自动化幼兔喂奶与清洗控制系统｜电控与嵌入式控制部分")
    style_run(r, size=13, color=MUTED)

    metadata = [
        ("负责方向", "单片机控制、双电机驱动、传感检测与系统调试"),
        ("控制核心", "STM32F407VET6"),
        ("当前阶段", "硬件资料分析、开发环境搭建、板卡烧录验证及控制方案确定"),
        ("总结日期", "2026年7月7日"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(label + "：")
        style_run(r1, bold=True)
        r2 = p.add_run(value)
        style_run(r2)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(10)
    rule.paragraph_format.space_after = Pt(12)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    add_callout(
        doc,
        "阶段结论",
        "目前已完成电控方案梳理、双路FOC控制板资料分析、开发与烧录环境搭建、ST-Link连接及原始固件写入验证，并初步确定“低速下降—电流突增判断接触—回抬10 mm—位置保持—执行喂奶”的控制路线。",
    )

    add_heading(doc, "一、项目概况与个人职责", 1)
    add_body(
        doc,
        "本项目拟设计一套基于单片机的自动化幼兔喂奶与清洗控制系统，通过电机、泵、电磁阀、编码器及相关传感器，实现工位移动、喂奶平台下压提示、自动供奶、残奶回吸、管路清洗和废液排放等功能。",
    )
    add_body(
        doc,
        "我目前主要负责电控与嵌入式控制部分，工作范围包括控制板与电机资料分析、开发环境搭建、程序烧录与通信验证、控制流程设计，以及后续电机、传感器、泵和电磁阀的协调控制。",
    )

    add_heading(doc, "二、目前已完成的主要任务", 1)

    add_heading(doc, "1. 系统电控方案梳理", 2)
    add_body(
        doc,
        "对项目总体功能进行了重新整理，明确控制核心为STM32F407VET6，所有阀门均采用电磁阀控制。系统电控对象包括双电机机构、接触检测、限位保护、奶泵、清水与清洁剂管路以及废液排放。",
    )
    add_body(
        doc,
        "已形成《自动化幼兔喂奶与清洗系统电控设计》思维导图，对控制核心、水平移动、下压机构、供奶回吸、清洗流程、电源与驱动关系进行了结构化整理。",
    )

    add_heading(doc, "2. 双路FOC控制板与电机资料分析", 2)
    add_body(
        doc,
        "阅读了双路FOC控制板的README、源程序和板卡接口说明，确认该板以STM32F407为核心，可同时驱动两路无刷电机，并提供电流采样、速度环、位置环和编码器反馈功能。原程序支持AS5600与AS5047两类磁编码器。",
    )
    add_body(
        doc,
        "确认原程序的串口控制方式包括开环、电流环、速度环和位置环；同时认识到该程序本质上是双电机FOC实验程序，不能直接等同于幼兔喂奶系统的完整控制程序，后续需要在其底层电机控制基础上加入项目专用状态机。",
    )

    add_heading(doc, "3. 开发环境与工具链搭建", 2)
    add_body(
        doc,
        "完成Keil MDK、STM32F4器件包、相关编译组件以及STM32CubeProgrammer的安装与配置。对原工程进行编译测试，确认原工程依赖ARM Compiler 5授权；切换编译器后又受到免费版代码体积限制，因此当前阶段采用现成固件直接烧录的方式完成板卡验证。",
    )
    add_body(
        doc,
        "完成CH340串口、STM32 Bootloader及ST-Link驱动识别过程，最终使用ST-Link通过SWD接口稳定连接STM32F407目标芯片。",
    )

    add_heading(doc, "4. 程序烧录与板卡验证", 2)
    add_body(
        doc,
        "先完成最小串口自检程序测试，确认STM32主控、程序下载和USB1串口通信链路能够正常工作。之后通过STM32CubeProgrammer执行芯片擦除，并将资料包中的原始F407FOCtest4.bin固件写入0x08000000地址。",
    )
    add_body(
        doc,
        "烧录完成后读取Flash首部数据，并与原始固件对应字节进行比对，结果一致，说明原始FOC固件已成功写入，板卡烧录链路已经打通。",
    )

    add_heading(doc, "5. 开机动画循环问题分析", 2)
    add_body(
        doc,
        "原始固件运行后，显示屏持续重复播放开机动画。通过检查main.c与AS5600.c，确认正常程序应只播放一次开机动画并进入主界面；当前循环现象属于程序在编码器初始化阶段反复复位。",
    )
    add_body(
        doc,
        "具体原因是原程序默认两路电机均使用AS5600编码器，当I2C读取失败时，程序会显示错误信息并调用系统复位。因此，该现象不是主控损坏或烧录失败，而是编码器未连接、编码器类型不匹配或I2C通信异常造成的。",
    )

    add_heading(doc, "6. 当前动作需求与控制路线确定", 2)
    add_body(
        doc,
        "目前已明确下压动作的目标：两个电机驱动机构低速下降，在接触到目标后停止继续下压，随后回抬约10 mm，并稳定保持该位置，为后续喂奶过程提供合适间距。",
    )
    add_body(
        doc,
        "接触判断拟优先利用FOC板现有电流采样功能。当机构接触目标后，电机负载与q轴电流会明显上升；程序对电流信号进行滤波，并要求其连续超过阈值一定时间后，才判定为有效接触，以减少启动冲击、导轨摩擦和机械卡顿造成的误判。",
    )

    add_heading(doc, "三、阶段任务完成情况", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    widths = [1700, 3000, 4660]
    set_table_widths(table, widths)
    headers = ["任务模块", "当前状态", "阶段结果"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        style_run(r, bold=True, color=DARK_BLUE)

    rows = [
        ("需求梳理", "已完成初步方案", "明确STM32F407VET6、双电机、电流触碰检测、回抬10 mm与位置保持需求"),
        ("资料整理", "已完成", "完成电控思维导图，并梳理电机、泵、电磁阀和传感器的控制关系"),
        ("开发环境", "基本完成", "Keil、器件包、CubeProgrammer及ST-Link驱动已安装，编译许可限制已识别"),
        ("板卡烧录", "已完成", "通过ST-Link成功写入原始FOC固件，并完成Flash数据比对验证"),
        ("故障分析", "已完成初步定位", "开机动画循环定位为编码器通信失败后程序主动复位"),
        ("应用程序", "待设计实现", "需在现有FOC底层上增加双电机同步、接触检测、回抬与保持状态机"),
    ]
    for row_data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i < 2 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            style_run(r, size=10.2)
    set_table_widths(table, widths)

    add_heading(doc, "四、拟采用的控制逻辑", 1)
    add_callout(
        doc,
        "动作流程",
        "系统初始化 → 编码器与限位检测 → 双电机低速下降 → 电流连续超阈值 → 停止下降 → 根据编码器回抬10 mm → 位置闭环保持 → 执行喂奶 → 完成后回到初始位置。",
    )
    add_body(
        doc,
        "程序建议采用状态机组织，主要状态包括INIT（初始化）、HOME（回零）、IDLE（等待）、DOWN（下降）、CONTACT（接触确认）、LIFT_10MM（回抬10 mm）、HOLD（位置保持）和FAULT（故障保护）。",
    )
    add_body(
        doc,
        "下降阶段使用低速速度环或连续更新的位置目标；检测到有效接触后，两路电机同时停止下行，并按照丝杆导程、传动比和编码器计数值换算10 mm回抬距离。回抬完成后切换至位置环保持，避免机构因重力或外力再次下沉。",
    )
    add_body(
        doc,
        "为提高安全性，电流检测不能作为唯一保护条件。后续应增加上限位，建议同时保留下限位或最大下降行程、动作超时、编码器失联、双电机位置差过大和过流保护。",
    )

    add_heading(doc, "五、当前存在的问题", 1)
    add_bullet(doc, "两台电机的具体型号、额定电压、电流、极对数及尾部编码器型号仍需最终确认。")
    add_bullet(doc, "需要确认两台电机是共同驱动同一升降平台，还是分别承担水平移动和升降动作。")
    add_bullet(doc, "回抬10 mm必须结合丝杆导程、减速比和编码器分辨率进行准确换算。")
    add_bullet(doc, "接触电流阈值需要在空载、正常摩擦、实际接触等条件下实验标定，并增加滤波、持续时间及迟滞。")
    add_bullet(doc, "原FOC工程存在Keil编译许可和体积限制，后续需要确定合法可持续的编译工具链。")
    add_bullet(doc, "泵、电磁阀和清洗管路尚未进入实际接线与联调阶段，后续应通过MOS管或继电器驱动模块与STM32控制信号连接。")

    add_heading(doc, "六、下一阶段工作计划", 1)
    add_number(doc, "确认电机型号、三相线定义、编码器型号和接口电压，正确配置AS5600或AS5047。")
    add_number(doc, "先进行单电机低速、速度环和位置环测试，再进行双电机同步测试，验证编码器方向与零位。")
    add_number(doc, "建立空载电流基线，记录下降、启动、机械摩擦和接触过程中的电流曲线，确定触碰阈值。")
    add_number(doc, "实现下降、接触确认、回抬10 mm和位置保持状态机，并加入超时、限位、过流和编码器失联保护。")
    add_number(doc, "完成双电机同步控制，监控两侧位置差；超过允许误差时立即停止并进入故障状态。")
    add_number(doc, "在电机控制稳定后，逐步接入奶泵、电磁阀、残奶回吸和清洗管路，完成整机流程联调。")

    add_heading(doc, "七、阶段性总结", 1)
    add_body(
        doc,
        "现阶段工作重点已从“识别板卡和搭建开发环境”推进到“确定项目专用控制方案”。目前已经证明STM32主控、串口、ST-Link和固件烧录链路可用，也已掌握原FOC程序的基本结构、编码器依赖及故障复位逻辑。",
    )
    add_body(
        doc,
        "下一阶段的核心任务是围绕真实机械机构完成电机和编码器验证，并在现有FOC底层基础上实现双电机下压、基于电流的接触检测、回抬10 mm和位置保持。完成该部分后，才能进一步接入喂奶、回吸和清洗执行部件。",
    )

    doc.core_properties.title = "自动化幼兔喂奶系统电控阶段性工作总结"
    doc.core_properties.subject = "电控与嵌入式控制阶段总结"
    doc.core_properties.author = "项目成员（电控方向）"
    doc.core_properties.keywords = "STM32F407, FOC, 双电机, 电流检测, 幼兔喂奶"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
